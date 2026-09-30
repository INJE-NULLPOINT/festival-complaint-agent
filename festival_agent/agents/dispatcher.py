"""③ 조치 에이전트 (Dispatcher) — Agent Path.

목표: 심각 이슈를 담당 부서가 바로 쓸 수 있는 문서로 만든다.

에이전트의 판단: 어떤 원문을 대표로 인용할지, 어떤 조치를 제안할지.
코드가 하는 일: 부서 매핑 조회, DOCX 생성, 상태 기록.
"""
import json
from datetime import datetime
from pathlib import Path

from core import config, db, issues, tourapi
from core.llm import Agent, tool

OUT_DIR = Path(__file__).resolve().parent.parent / "output"

# 파일명에 쓸 수 없는 문자. 라벨에 '/'가 들어있어("안내/동선") 경로로 해석되면서
# DOCX 저장이 실패한 적이 있다.
_BAD_CHARS = str.maketrans({ch: "_" for ch in '\\/:*?"<>|\r\n\t'})


def _safe(name: str) -> str:
    """Windows 파일명으로 안전한 문자열로 바꾼다."""
    out = (name or "").translate(_BAD_CHARS).strip(" .")
    return out or "미상"

SYSTEM = """너는 지역 축제 운영 관제 시스템의 '조치 에이전트'다.

역할: 심각도가 높은 민원 유형에 대해, 담당 부서가 받아서 바로 움직일 수 있는
조치요청서를 만든다.

절차
1. get_department 로 담당 부서를 확인한다.
2. collect_quotes 로 해당 유형의 민원 원문을 가져온다.
3. 그중 **대표성 있는 5건 이내**를 직접 고른다. 고르는 기준:
   - 서로 다른 구역의 사례를 섞는다 (한 구역에 몰리지 않게)
   - 구체적 상황이 드러난 문장을 고른다 ("불편함" 같은 모호한 것은 제외)
   - 안전 관련이면 위험 상황이 명확한 것을 우선한다
4. lookup_festival_info 로 축제 공식 정보를 조회해 문서 머리말에 넣는다.
   조회에 실패하면(null) 축제명만 쓰고 넘어간다. 없는 정보를 지어내지 마라.
5. 조치 제안을 쓴다. **현장에서 오늘 실행 가능한 것만** 쓴다.
   예산 편성, 조례 개정 같은 장기 과제는 쓰지 않는다.
   요청문에 '관제 카드의 조치'가 주어지면 그 문장을 그대로(문장·순서 유지) 제안으로 쓴다.
   관제 화면과 요청서의 조치가 서로 다르면 안 된다. 주어지지 않았을 때만 2~3개를 직접 쓴다.
6. generate_doc 으로 문서를 만들고, 무엇을 만들었는지 한 문장으로 보고한다.

반드시 지킬 것
- 인용은 원문 그대로 쓴다. 문장을 고쳐 쓰지 마라.
- 건수와 점수는 주어진 값을 그대로 쓴다. 추정하지 마라.
"""


@tool(
    name="get_department",
    description="민원 유형의 담당 부서와 연락처를 조회한다.",
    properties={"label": {"type": "string", "enum": list(config.LABELS)}},
    required=["label"],
)
def get_department(label: str) -> dict:
    dept, contact = config.DEPARTMENT_MAP.get(label, ("미지정", "-"))
    return {"label": label, "korean": config.LABELS.get(label, label),
            "department": dept, "contact": contact}


@tool(
    name="collect_quotes",
    description="해당 유형의 민원 원문을 최근 순으로 가져온다. 인용할 후보다.",
    properties={
        "label": {"type": "string", "enum": list(config.LABELS)},
        "limit": {"type": "integer", "description": "최대 건수 (기본 15)"},
    },
    required=["label"],
)
def collect_quotes(label: str, limit: int = 15) -> list[dict]:
    with db.connect() as conn:
        rows = conn.execute(
            f"""SELECT f.raw_text, COALESCE(z.name, '{config.ZONE_UNKNOWN}') zone, f.ingested_at,
                      c.sentiment, c.confidence
               FROM classification c
               JOIN feedback f ON f.id = c.feedback_id
               LEFT JOIN zone z ON z.id = f.zone_id
               WHERE c.label=? AND c.status='done' AND f.deleted_at IS NULL
               ORDER BY f.id DESC LIMIT ?""",
            (label, limit),
        ).fetchall()
    return [dict(r) for r in rows]


@tool(
    name="generate_doc",
    description="부서별 조치요청서 DOCX 파일을 생성한다.",
    properties={
        "label": {"type": "string", "enum": list(config.LABELS)},
        "department": {"type": "string"},
        "count": {"type": "integer", "description": "해당 유형 민원 건수"},
        "score": {"type": "number", "description": "심각도 점수"},
        "grade": {"type": "string"},
        "formula": {"type": "string", "description": "심각도 계산식 (근거로 문서에 싣는다)"},
        "quotes": {
            "type": "array",
            "description": "인용할 민원 원문 (최대 5건)",
            "items": {
                "type": "object",
                "properties": {
                    "raw_text": {"type": "string"},
                    "zone": {"type": "string"},
                    "ingested_at": {"type": "string"},
                },
                "required": ["raw_text"],
                "additionalProperties": False,
            },
        },
        "suggestions": {
            "type": "array",
            "description": "즉시 실행 가능한 조치 제안 2~4개 (관제 카드의 조치가 있으면 그대로)",
            "items": {"type": "string"},
        },
        "festival_info": {"type": "string", "description": "축제 공식 정보 한 줄. 없으면 빈 문자열"},
    },
    required=["label", "department", "count", "score", "grade", "quotes", "suggestions"],
)
def generate_doc(label: str, department: str, count: int, score: float,
                 grade: str, quotes: list, suggestions: list,
                 formula: str = "", festival_info: str = "") -> dict:
    from docx import Document
    from docx.shared import Pt

    OUT_DIR.mkdir(exist_ok=True)
    korean = config.LABELS.get(label, label)
    grade_ko = {"immediate": "즉시 조치", "high": "높음",
                "mid": "보통", "low": "낮음"}.get(grade, grade)

    doc = Document()
    doc.add_heading("축제 민원 조치요청서", level=0)

    with db.connect() as conn:
        fest = conn.execute("SELECT name FROM festival LIMIT 1").fetchone()
    head = doc.add_paragraph()
    head.add_run(f"{fest['name'] if fest else ''}\n").bold = True
    head.add_run(f"수신: {department}\n")
    head.add_run(f"작성: 실시간 민원 관제 AI Agent\n")
    head.add_run(f"일시: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n")
    if festival_info:
        head.add_run(f"{festival_info}\n").italic = True

    doc.add_heading("1. 요청 사유", level=1)
    t = doc.add_table(rows=4, cols=2)
    t.style = "Light Grid Accent 1"
    for i, (k, v) in enumerate([
        ("민원 유형", korean),
        ("접수 건수", f"{count}건"),
        ("심각도", f"{score}점 ({grade_ko})"),
        ("판정 근거", formula or "-"),
    ]):
        t.rows[i].cells[0].text = k
        t.rows[i].cells[1].text = str(v)

    doc.add_heading("2. 접수된 민원 (원문 인용)", level=1)
    for q in quotes[:5]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"“{q.get('raw_text','')}”")
        meta = f"  — {q.get('zone') or '구역 미상'}"
        if q.get("ingested_at"):
            meta += f", {str(q['ingested_at'])[11:16]}"
        r = p.add_run(meta)
        r.font.size = Pt(9)

    doc.add_heading("3. 조치 제안", level=1)
    for s in suggestions[:4]:
        doc.add_paragraph(s, style="List Number")

    doc.add_heading("4. 비고", level=1)
    doc.add_paragraph(
        "본 문서는 실시간 민원 관제 AI Agent가 자동 생성했습니다. "
        "심각도는 검증된 계산 함수가 산출하며, 동일 입력에 동일 결과가 나옵니다. "
        "민원 원문에서 개인정보는 접수 시점에 자동 마스킹되었습니다."
    )

    # 라벨에 '/'가 들어간다("안내/동선"). 파일명에 쓰면 경로로 해석되므로 정리한다.
    stamp = datetime.now().strftime("%m%d_%H%M%S")
    path = OUT_DIR / f"조치요청서_{_safe(department)}_{_safe(korean)}_{stamp}.docx"
    doc.save(path)

    # 웹 화면이 DOCX 를 열지 않고도 같은 내용을 미리 볼 수 있게 한 벌 더 남긴다.
    preview = {
        "festival": fest["name"] if fest else "",
        "department": department,
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "festival_info": festival_info,
        "label": label, "label_ko": korean, "count": count,
        "score": score, "grade": grade, "grade_ko": grade_ko, "formula": formula,
        "quotes": [{"raw_text": q.get("raw_text", ""), "zone": q.get("zone") or "",
                    "time": str(q.get("ingested_at") or "")[11:16]} for q in quotes[:5]],
        "suggestions": list(suggestions[:4]),
    }
    url = _upload(path, f"{label}_{stamp}.docx")

    with db.connect() as conn:
        # 같은 유형의 열린 이전 요청서는 새 요청서로 대체된다 (완료 건은 그대로)
        conn.execute(
            f"UPDATE action_request SET status='superseded', closed_at=? "
            f"WHERE label=? AND {db.OPEN_ACTION_SQL}",
            (db.now(), label),
        )
        cur = conn.execute(
            """INSERT INTO action_request
               (festival_id, label, department, count, doc_path, status, created_at,
                doc_url, doc_json)
               VALUES (?,?,?,?,?, 'requested', ?, ?, ?)""",
            (db.festival_id(), label, department, count, str(path), db.now(),
             url, json.dumps(preview, ensure_ascii=False)),
        )
        conn.commit()
        action_id = cur.lastrowid

    return {"action_id": action_id, "path": str(path),
            "quotes_used": len(quotes[:5])}


def _upload(path: Path, key: str) -> str | None:
    """Supabase Storage 에 올리고 공개 URL 을 돌려준다.

    Storage 설정이 없으면 local 대역(webapi.py)이 output/ 에서 서빙하는 URL 을
    돌려준다. 업로드가 실패하면 None — 문서 생성을 막지는 않는다.
    """
    if not (config.SUPABASE_URL and config.SUPABASE_SERVICE_KEY):
        from urllib.parse import quote
        return f"/api/docs/{quote(path.name)}"
    import urllib.request
    # Storage 키는 ASCII 로 둔다. 한글 파일명은 인코딩 문제로 실패하는 경우가 있다.
    obj = f"{config.SUPABASE_BUCKET}/{key}"
    req = urllib.request.Request(
        f"{config.SUPABASE_URL}/storage/v1/object/{obj}",
        data=path.read_bytes(), method="POST",
        headers={
            "Authorization": f"Bearer {config.SUPABASE_SERVICE_KEY}",
            "apikey": config.SUPABASE_SERVICE_KEY,
            "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "x-upsert": "true",
        },
    )
    try:
        urllib.request.urlopen(req, timeout=15).read()
    except Exception as e:
        db.log_agent("dispatcher", "upload_failed", output_summary=str(e))
        return None
    return f"{config.SUPABASE_URL}/storage/v1/object/public/{obj}"


def _festival_tool():
    from .common import lookup_festival_info
    return lookup_festival_info


def local_run(agent, user_input: str, ctx: dict) -> str:
    """local 대역 — 구역이 겹치지 않게 인용을 고르고 템플릿 제안을 쓴다. 제출본 아님."""
    from core import rules

    label = ctx["label"]
    dept = agent.call("get_department", label=label)
    quotes = agent.call("collect_quotes", label=label, limit=15)
    picked = rules.pick_quotes(quotes, n=5)

    info = agent.call("lookup_festival_info", keyword=ctx.get("festival_keyword", ""))
    line = ""
    if info:
        line = f"{info.get('title','')} · {info.get('period','')} · {info.get('addr','')}"

    res = agent.call(
        "generate_doc",
        label=label, department=dept["department"], count=ctx["count"],
        score=ctx["score"], grade=ctx["grade"], formula=ctx.get("formula", ""),
        quotes=picked, suggestions=ctx.get("card_actions") or rules.suggestions_for(label),
        festival_info=line,
    )
    return (f"{dept['department']} 조치요청서 생성 "
            f"(인용 {res['quotes_used']}건) — {res['path']}  (local 대역)")


dispatcher = Agent(
    name="dispatcher",
    system=SYSTEM,
    tools=[get_department, collect_quotes, _festival_tool(), generate_doc],
    max_steps=10,
    max_tokens=8192,
    local=local_run,
)


def run_for(label: str, score: float, grade: str, formula: str = "",
            window_min: int = config.DEFAULT_WINDOW_MIN) -> str:
    """특정 유형에 대해 조치요청서를 만든다.

    건수는 심각도를 판정한 창과 같은 구간에서 센다. 창이 다르면 문서의 건수와
    심각도 근거가 어긋난다.
    """
    korean = config.LABELS.get(label, label)
    cnt = db.label_counts(window_min).get(label, 0)
    with db.connect() as conn:
        fest = conn.execute("SELECT name FROM festival LIMIT 1").fetchone()
    # 건수·축제명은 모델이 도구로 알아낼 수 없다. 요청문에 직접 넣어야
    # generate_doc 의 count 를 채운다 (빠지면 모델이 문서 생성을 멈춘다).
    fest_name = fest["name"] if fest else ""
    # 관제 카드에서 AI 가 근거 민원을 읽고 정리한 조치 — 요청서와 관제 화면의 조치가 같게 넘긴다
    card_actions = issues.actions_for_label(label)
    card_line = ("\n관제 카드의 조치(제안에 이 문장을 그대로 써라): "
                 + " / ".join(card_actions)) if card_actions else ""
    return dispatcher.run(
        f"'{korean}'({label}) 유형의 심각도가 {score}점({grade})으로 판정됐다. "
        f"판정 근거는 다음과 같다: {formula}\n"
        f"같은 구간(최근 {window_min}분)의 접수 건수는 {cnt}건이다. "
        f"축제: {fest_name or '미상'}\n"
        f"담당 부서용 조치요청서를 만들어줘.{card_line}",
        ctx={"label": label, "score": score, "grade": grade, "formula": formula,
             "count": cnt, "festival_keyword": (fest["name"] if fest else "")[:4],
             "card_actions": card_actions},
    )


def pending_labels(min_grade: tuple[str, ...] = ("immediate", "high")) -> list[dict]:
    """조치요청서가 아직 없는 심각 유형을 찾는다."""
    ranked = db.ranked()
    out = []
    with db.connect() as conn:
        for r in ranked:
            if r["grade"] not in min_grade:
                continue
            exists = conn.execute(
                f"SELECT id FROM action_request WHERE label=? AND {db.OPEN_ACTION_SQL}",
                (r["label"],),
            ).fetchone()
            if not exists:
                out.append(r)
    return out
