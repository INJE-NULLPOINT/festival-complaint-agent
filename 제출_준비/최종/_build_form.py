"""운영규정 docx의 별지 1 표에 기술설명서.md 내용을 채워 신청서(최종수정본)를 만든다.

    python 제출_준비/최종/_build_form.py
"""
import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Mm, Pt
from docx.table import Table

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _build_docs import to_pdf  # noqa: E402

HERE = Path(__file__).resolve().parent
FORM = Path(r"C:\Users\12kjm\Downloads\03-1_제4회_경남AI_SW경진대회_운영규정v1.0_최종_20260907.docx")
OUT = HERE / "신청서_최종수정본_철철철.docx"
FONT_PT = 9  # 양식 9.5pt → 9pt (1쪽에 맞춤)
MARGIN_MM = 10  # 양식 위아래 16mm → 10mm
CELL_PAD = 30  # 칸 안 위아래 여백 양식 80 → 30 (dxa)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# 1쪽에 맞추려고 신청서에서만 줄여 쓴 칸 (숫자·사실은 기술설명서와 같음). 키 = 표 행 번호(머리줄 0)
SHORT = {
    1: ["철철철 / 대학부 / ⑤ 경남 지역혁신·공공서비스 AI Agent"],
    3: ["건수 순이면 적은 안전 민원(어두운 길·난간·밀림)이 불편 민원(주차·가격)에 묻힌다"],
    4: ["축제 운영 담당 주무관·현장 운영요원(관제·조치 화면), 방문객(QR 접수)"],
    5: ["민원을 분류해 심각도 순으로 \"지금 조치할 일\"과 담당 부서 조치요청서를 만든다"],
    6: ["Claude Opus 5.5(claude-opus-5-5), Claude Code CLI 경유. 점수는 결정적 함수(severity.ts)"],
    7: ["도구 17종(예: plan_cycle, score_label, generate_doc, write_briefing)",
        "Data: 합성 시드 160건, 축제 정보 수기(TourAPI 미사용), 이태원 112 녹취 11건. DB: Supabase/SQLite, DOCX"],
    8: ["- State: 민원·분류·카드·조치 상태·실행 기록(agent_log)을 DB에 유지",
        "- Memory: 비슷한 과거 민원 참고(lookup_similar), 운영자가 정한 유형 우선",
        "- Feedback: 조치 상태 → 다음 판정(S-06 미조치 30분 ×1.2), '확인 필요' 지정 → 기억"],
    9: ["① 방문객이 QR 화면에서 민원 접수",
        "② 워커가 개인정보를 가려 저장, 접수 원문 삭제",
        "③ ①분류: 유형·안전 여부(신뢰도 0.3 미만은 '확인 필요')",
        "④ ⓪계획이 단계 결정, ②감시가 심각도 함수 호출로 등급",
        "⑤ ③조치: 담당 부서 조치요청서(DOCX)",
        "⑥ ④통합: 관제 카드·브리핑",
        "⑦ 운영자의 조치 상태 변경 → 다음 주기 판정에 반영"],
    11: ["- 기존자산: 자체 기존 코드 없음. 외부: Node.js·TypeScript·Vite·Supabase 등 오픈소스(저장소 package.json), "
         "Claude(Anthropic), 행안부 매뉴얼·인파 가이드라인, 이태원 112 공개 녹취",
         "- 8일 신규개발: 에이전트 5종·루프·작업 큐(직접 구현), 심각도 함수, 계획·기억, 관제 카드, "
         "마스킹·인젝션 탐지, 부서 매핑·DOCX, 리플레이·합성 시드, 측정 도구, 웹 앱"],
    12: ["5/5 통과(정상·모호·빈 입력·API 오류·악의적 입력, 실제 모델) · 공격 점검 26/26",
         "정확도 96.9%(31/32, 합성 참고값) · 112 녹취 재현 18:34 첫 신고 '즉시'(규칙)"],
    15: ["행안부 지역축제장 안전관리 매뉴얼·다중운집인파사고 가이드라인(2024)",
         "민원처리법 시행령 제19조 · Claude Code(작성 보조) · 상세는 출처신고서"],
}


def rows_from_md():
    src = (HERE.parent / "기술설명서.md").read_text(encoding="utf-8")
    src = re.sub(r"<!--[\s\S]*?-->", "", src)
    rows = [l.strip().strip("|").split("|") for l in src.splitlines() if l.startswith("|") and not re.fullmatch(r"\|[\s|:-]+\|", l.strip())]
    out = []
    for r in rows[1:]:
        v = re.sub(r"`([^`]*)`", r"\1", r[1].strip()).replace("**", "")
        out.append([s.strip() for s in re.split(r" / (?=[-①②③④⑤⑥⑦])|\s(?=[2-5]\) )|\s{2,}·\s|(?<=\))\s{2,}", v) if s.strip()])
    return out


def fill(cell, lines, rpr):
    ps = cell.paragraphs
    first = ps[0]
    for p in ps[1:]:
        p._element.getparent().remove(p._element)
    for r in list(first._element.findall(W + "r")):
        first._element.remove(r)
    for k, text in enumerate(lines):
        p = first if k == 0 else cell.add_paragraph()
        if k:
            p._element.getparent().remove(p._element)
            p._element = copy.deepcopy(first._element)
            for r in list(p._element.findall(W + "r")):
                p._element.remove(r)
            cell._element.append(p._element)
        run = p._element.makeelement(W + "r", {})
        run.append(copy.deepcopy(rpr))
        t = run.makeelement(W + "t", {"{http://www.w3.org/XML/1998/namespace}space": "preserve"})
        t.text = text
        run.append(t)
        p._element.append(run)


def main():
    d = Document(FORM)
    body = d.element.body
    kids = list(body.iterchildren())
    title_i = next(i for i, el in enumerate(kids) if el.tag == W + "p" and "<별지 1>" in "".join(x.text or "" for x in el.iter(W + "t")))
    keep = {kids[title_i], kids[title_i + 1]}
    for el in kids:
        if el not in keep and el.tag != W + "sectPr":
            body.remove(el)
    t = Table(kids[title_i + 1], d)
    rpr = t.rows[2].cells[1].paragraphs[0].runs[0]._element.find(W + "rPr")
    head = t.rows[0].cells[1]
    hrpr = head.paragraphs[0].runs[0]._element.find(W + "rPr")
    fill(head, ["기재 내용"], hrpr)
    data = rows_from_md()
    assert len(data) == len(t.rows) - 1, (len(data), len(t.rows))
    for i, (row, lines) in enumerate(zip(t.rows[1:], data), 1):
        fill(row.cells[1], SHORT.get(i, lines), rpr)
    for row in t.rows:  # 샘플 문구에 맞춘 최소 높이(32·55mm)를 풀어 내용 높이로
        h = row._tr.find(f".//{W}trHeight")
        if h is not None and int(h.get(W + "val")) > 1000:
            h.getparent().remove(h)
    for mar in t._tbl.iter(W + "tcMar"):  # 칸 안 위아래 여백 80 → CELL_PAD (dxa)
        for side in ("top", "bottom"):
            el = mar.find(W + side)
            if el is not None:
                el.set(W + "w", str(CELL_PAD))
    for row in t.rows if FONT_PT else []:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(FONT_PT)
    tail = t._tbl.makeelement(W + "p", {})  # 표 뒤 필수 빈 문단이 2쪽으로 밀리지 않게 1pt
    tail.append(tail.makeelement(W + "pPr", {}))
    sp = tail[0].makeelement(W + "spacing", {W + "after": "0", W + "line": "20", W + "lineRule": "exact"})
    tail[0].append(sp)
    t._tbl.addnext(tail)
    sec = d.sections[0]
    for ref in sec._sectPr.findall(W + "headerReference") + sec._sectPr.findall(W + "footerReference"):  # 머리글 "운영규정(안)"·쪽 번호 삭제
        sec._sectPr.remove(ref)
    sec.top_margin = sec.bottom_margin = Mm(MARGIN_MM)
    d.save(OUT)
    to_pdf([OUT])


if __name__ == "__main__":
    main()
